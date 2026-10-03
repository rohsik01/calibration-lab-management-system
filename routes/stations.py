"""Route module: stations."""
from app import *

@app.route("/register")
def register():
    rows = get_db().execute(LATEST + " ORDER BY s.sensor_id").fetchall()
    return render_template("register.html", rows=rows)



def _delete_station_records(db, station_ids):
    station_ids = [int(x) for x in station_ids]
    if not station_ids:
        return 0
    sp = ",".join("?" * len(station_ids))
    sensor_rows = db.execute(f"SELECT sensor_id FROM sensors WHERE station_id IN ({sp})", station_ids).fetchall()
    sensor_ids = [r["sensor_id"] for r in sensor_rows]
    if sensor_ids:
        xp = ",".join("?" * len(sensor_ids))
        cal_rows = db.execute(f"SELECT cal_id FROM calibrations WHERE sensor_id IN ({xp})", sensor_ids).fetchall()
        cal_ids = [r["cal_id"] for r in cal_rows]
        if cal_ids:
            cp = ",".join("?" * len(cal_ids))
            db.execute(f"DELETE FROM calibration_points WHERE cal_id IN ({cp})", cal_ids)
            db.execute(f"DELETE FROM calibrations WHERE cal_id IN ({cp})", cal_ids)
        db.execute(f"DELETE FROM sensors WHERE sensor_id IN ({xp})", sensor_ids)
    db.execute(f"DELETE FROM stations WHERE station_id IN ({sp})", station_ids)
    return len(station_ids)


@app.route("/stations/<int:station_id>/delete", methods=["POST"])
@admin_required
def delete_station(station_id):
    db = get_db()
    station = db.execute("SELECT station_id, name FROM stations WHERE station_id=?", (station_id,)).fetchone()
    if not station:
        abort(404)
    try:
        with db:
            _delete_station_records(db, [station_id])
        flash(f"Station '{station['name']}' and its sensors/calibration records were deleted.")
    except sqlite3.Error:
        flash("Could not delete the station and its related records.", "error")
    return redirect(url_for("stations"))


@app.route("/stations/bulk-delete", methods=["POST"])
@admin_required
def bulk_delete_stations():
    raw_ids = request.form.getlist("station_ids")
    station_ids = []
    for value in raw_ids:
        try:
            station_ids.append(int(value))
        except (TypeError, ValueError):
            continue
    station_ids = list(dict.fromkeys(station_ids))
    if not station_ids:
        flash("Select at least one station to delete.", "error")
        return redirect(url_for("stations"))
    db = get_db()
    placeholders = ",".join("?" * len(station_ids))
    existing = db.execute(
        f"SELECT station_id FROM stations WHERE station_id IN ({placeholders})", station_ids
    ).fetchall()
    existing_ids = [r["station_id"] for r in existing]
    try:
        with db:
            deleted = _delete_station_records(db, existing_ids)
        flash(f"{deleted} station(s) and their sensors/calibration records were deleted.")
    except sqlite3.Error:
        flash("Could not delete the selected stations and their related records.", "error")
    return redirect(url_for("stations"))


@app.route("/stations", methods=["GET", "POST"])
@admin_required
def stations():
    db = get_db()
    if request.method == "POST":
        try:
            station_id_raw = request.form.get("station_id", "").strip()
            name = request.form["name"].strip()
            location = request.form.get("location", "").strip()
            station_type = request.form.get("type", "").strip() or "Meteorological"
            if not station_id_raw or not station_id_raw.isdigit() or int(station_id_raw) <= 0:
                raise ValueError("A valid positive Station ID is required.")
            station_id = int(station_id_raw)
            if not name:
                raise ValueError("Station name is required.")
            if db.execute("SELECT 1 FROM stations WHERE station_id=?", (station_id,)).fetchone():
                raise ValueError(f"Station ID {station_id} is already in use.")
            db.execute("INSERT INTO stations(station_id, name, location, type, updated_at) VALUES (?,?,?,?,?)",
                       (station_id, name, location, station_type, datetime.now().isoformat(timespec="seconds")))
            db.commit()
            flash("Station added.")
        except sqlite3.IntegrityError:
            flash("That station already exists.")
        except ValueError as e:
            flash(str(e), "error")
        return redirect(url_for("stations"))
    return render_template("stations.html",
                           rows=db.execute("SELECT * FROM stations ORDER BY updated_at DESC, station_id DESC").fetchall())


def _excel_workbook(instructions, sheet_name, headers, sample_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    info = wb.create_sheet("Instructions")
    info.append(["Calibration Lab Management System - Bulk Upload"])
    for line in instructions:
        info.append([line])
    info["A1"].font = Font(bold=True, size=14)
    for cell in info["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 110
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5FBF")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in sample_rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@app.route("/stations/bulk-sample")
def station_bulk_sample():
    data = _excel_workbook(
        [
            "Fill one station per row in the 'Stations' sheet.",
            "Station ID: enter the unique positive numeric Station ID assigned to the station. It is required for every station record.",
            "Station Name is required and must be unique.",
            "Location and Type are required for complete station details. Example Type: Meteorological, Hydrological, Agrometeorological, Radar.",
            "Do not change the column headings."
        ],
        "Stations",
        ["Station ID", "Station Name", "Location", "Type"],
        [[1001, "Example Station", "Dharan, Sunsari", "Meteorological"]]
    )
    return Response(data, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="station_bulk_upload_sample.xlsx"'})


@app.route("/stations/bulk-upload", methods=["POST"])
@admin_required
def station_bulk_upload():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("Please choose an Excel .xlsx file.", "error")
        return redirect(url_for("stations"))
    try:
        from openpyxl import load_workbook
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb["Stations"] if "Stations" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("The Excel file is empty.")
        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        expected = ["Station ID", "Station Name", "Location", "Type"]
        if headers != expected:
            raise ValueError("Invalid columns. Download the station Excel sample and use its column headings.")
        db = get_db()
        parsed = []
        errors = []
        seen_ids = set()
        seen_names = set()
        for row_no, row in enumerate(rows[1:], 2):
            if not any(v not in (None, "") for v in row):
                continue
            vals = list(row) + [""] * (4 - len(row))
            sid, name, location, station_type = vals[:4]
            name = str(name).strip() if name is not None else ""
            location = str(location).strip() if location is not None else ""
            station_type = str(station_type).strip() if station_type is not None else ""
            if not name:
                errors.append(f"Row {row_no}: Station Name is required.")
                continue
            if not station_type:
                errors.append(f"Row {row_no}: Type is required.")
                continue
            sid_int = None
            if sid in (None, ""):
                errors.append(f"Row {row_no}: Station ID is required.")
                continue
            try:
                sid_int = int(float(sid))
                if sid_int <= 0:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append(f"Row {row_no}: Station ID must be a positive number.")
                continue
            if sid_int in seen_ids:
                errors.append(f"Row {row_no}: duplicate Station ID {sid_int} in the file.")
                continue
            seen_ids.add(sid_int)
            name_key = name.casefold()
            if name_key in seen_names:
                errors.append(f"Row {row_no}: duplicate Station Name '{name}'.")
                continue
            seen_names.add(name_key)
            parsed.append((sid_int, name, location, station_type, row_no))
        if errors:
            raise ValueError("Upload stopped. " + " ".join(errors[:12]) +
                             (f" Showing first 12 of {len(errors)} errors." if len(errors) > 12 else ""))
        with db:
            for sid, name, location, station_type, row_no in parsed:
                if sid is None:
                    raise ValueError(f"Row {row_no}: Station ID is required.")
                else:
                    existing = db.execute("SELECT station_id FROM stations WHERE station_id=?", (sid,)).fetchone()
                    if existing:
                        conflict = db.execute("SELECT station_id FROM stations WHERE name=? COLLATE NOCASE AND station_id<>?",
                                               (name, sid)).fetchone()
                        if conflict:
                            raise ValueError(f"Row {row_no}: station name '{name}' belongs to another station.")
                        db.execute("UPDATE stations SET name=?, location=?, type=?, updated_at=? WHERE station_id=?",
                                   (name, location, station_type, datetime.now().isoformat(timespec="seconds"), sid))
                    else:
                        if db.execute("SELECT 1 FROM stations WHERE name=? COLLATE NOCASE", (name,)).fetchone():
                            raise ValueError(f"Row {row_no}: station name '{name}' already exists.")
                        db.execute("INSERT INTO stations(station_id, name, location, type) VALUES (?,?,?,?)",
                                   (sid, name, location, station_type))
        flash(f"{len(parsed)} station record(s) uploaded successfully.")
    except ValueError as e:
        flash(str(e), "error")
    except Exception as e:
        flash(f"Could not process the Excel file: {e}", "error")
    return redirect(url_for("stations"))




