"""Route module: reports."""
from app import *

# Export/report definitions shared by the report routes.
EXPORT_COLUMNS = [
    ("sensor_id", "Sensor ID"), ("station", "Station"), ("sensor_type", "Sensor type"),
    ("manufacturer", "Manufacturer"), ("serial_number", "Serial number"),
    ("cal_date", "Calibration date"), ("reference_standard", "Reference standard"),
    ("n_points", "Number of points"), ("point_no", "Point no."),
    ("point_tolerance", "Point tolerance (+/-)"),
    ("reference_value", "Reference value"), ("measured_value", "Results (reading)"),
    ("error", "Error"), ("result", "Pass/Fail"), ("overall_result", "Overall result"),
    ("certificate_no", "Certificate"), ("next_due", "Next due date"),
    ("performed_by", "Calibrated by"), ("unit", "Unit"), ("tolerance", "Tolerance (+/-)"),
    ("status", "Status"),
]

HISTORY_SQL = """
SELECT s.*, st.name AS station, c.cal_date, c.reference_standard, c.reference_value,
       c.measured_value, c.error, c.result, c.certificate_no, c.next_due, c.performed_by,
       c.n_points
FROM calibrations c JOIN sensors s USING(sensor_id) JOIN stations st USING(station_id)
"""

POINTS_SQL = """
SELECT s.*, st.name AS station, c.cal_date, c.reference_standard, c.n_points, p.point_no,
       p.reference_value, p.measured_value, p.error, p.result, p.tolerance AS point_tolerance,
       c.result AS overall_result,
       c.certificate_no, c.next_due, c.performed_by
FROM calibration_points p JOIN calibrations c USING(cal_id)
JOIN sensors s USING(sensor_id) JOIN stations st USING(station_id)
"""

STATUS_FILTERS = {"ok": "OK", "overdue": "Overdue", "failed": "Failed",
                  "never": "Never calibrated", "due_soon": "Due in"}


def export_rows(args):
    """Rows for the export, filtered by the query-string options. Returns (rows, scope)."""
    scope = args.get("scope") if args.get("scope") in ("history", "points") else "latest"
    where, params = [], []
    stations_ = [int(x) for x in args.getlist("station") if x.isdigit()]
    if stations_:
        where.append(f"s.station_id IN ({','.join('?' * len(stations_))})")
        params += stations_
    sensors_ = [x for x in args.getlist("sensor") if x]
    if sensors_:
        where.append(f"s.sensor_id IN ({','.join('?' * len(sensors_))})")
        params += sensors_
    for key, op in (("date_from", ">="), ("date_to", "<=")):
        val = args.get(key, "")
        try:
            date.fromisoformat(val)
        except ValueError:
            continue
        where.append(f"c.cal_date {op} ?")
        params.append(val)
    if args.get("result") in ("PASS", "FAIL"):
        where.append(("p.result" if scope == "points" else "c.result") + " = ?")
        params.append(args["result"])
    sql = {"history": HISTORY_SQL, "points": POINTS_SQL}.get(scope, LATEST)
    sql += (" WHERE " + " AND ".join(where) if where else "")
    sql += " ORDER BY s.sensor_id" + (", c.cal_date, c.cal_id" if scope != "latest" else "")
    sql += ", p.point_no" if scope == "points" else ""
    rows = [dict(r) for r in get_db().execute(sql, params)]
    for r in rows:
        r["status"] = status(r)[0] if scope == "latest" else ""
        r.setdefault("point_no", "")
        r.setdefault("point_tolerance", "")
        r.setdefault("overall_result", r.get("result"))
    wanted = STATUS_FILTERS.get(args.get("status", ""))
    if wanted and scope == "latest":
        rows = [r for r in rows if r["status"].startswith(wanted)]
    return rows, scope


def safe_cell(v):
    """Stop spreadsheet programs from running text that starts like a formula."""
    if v is None:
        return ""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


@app.route("/export")
def export_csv():
    db = get_db()
    return render_template(
        "export.html", columns=EXPORT_COLUMNS,
        stations=db.execute("SELECT * FROM stations ORDER BY name").fetchall(),
        sensors=db.execute("SELECT s.sensor_id, s.sensor_type, s.station_id, st.name AS station "
                           "FROM sensors s JOIN stations st USING(station_id) "
                           "ORDER BY s.sensor_id").fetchall())


DATE_KEYS = {"cal_date", "next_due", "calibrated_on", "valid_until"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STD_SHEET_COLUMNS = [("code", "Code"), ("name", "Name"), ("standard_type", "Type"),
                     ("manufacturer", "Manufacturer"), ("serial_number", "Serial number"),
                     ("uncertainty", "Uncertainty"), ("traceability", "Traceability"),
                     ("certificate_no", "Certificate"), ("calibrated_on", "Calibrated on"),
                     ("valid_until", "Valid until"), ("status", "Status"), ("active_txt", "Active")]


def export_label(args):
    sensors_ = [x for x in args.getlist("sensor") if x]
    stations_ = [x for x in args.getlist("station") if x.isdigit()]
    if len(sensors_) == 1:
        label = sensors_[0]
    elif len(stations_) == 1:
        st = get_db().execute("SELECT name FROM stations WHERE station_id=?",
                              (int(stations_[0]),)).fetchone()
        label = st["name"] if st else "station"
    else:
        label = "selection" if (sensors_ or stations_) else "all"
    return re.sub(r"[^A-Za-z0-9_-]+", "-", label).strip("-") or "export"


def export_value(key, v):
    """Pass/Fail and status words follow the report language."""
    if g.get("lang") != "ne" or v in (None, ""):
        return v
    if key in ("result", "overall_result"):
        return tr(v)
    if key == "status":
        return status_label(v)
    return v


def build_xlsx(rows, cols, scope, per_station, with_standards, banner):
    """Excel workbook with the banner on top of every sheet. Raises ImportError without openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fills = {k: PatternFill("solid", fgColor=v) for k, v in
             dict(green="D5F0DD", red="F8D4D1", amber="FBE8C4", grey="E3E6EB").items()}
    head_fill = PatternFill("solid", fgColor="14203A")
    wb = Workbook()
    wb.remove(wb.active)
    used = set()

    def new_sheet(name):
        base = re.sub(r"[\[\]:*?/\\]", "-", name)[:28] or "Sheet"
        n, i = base, 2
        while n.lower() in used:
            n, i = f"{base[:25]}-{i}", i + 1
        used.add(n.lower())
        return wb.create_sheet(n)

    def color(key, v):
        if key in ("result", "overall_result"):
            return "green" if v == "PASS" else "red" if v == "FAIL" else None
        if key == "status" and v:
            v = str(v)
            return ("green" if v in ("OK", "Valid") else "amber" if v.startswith(("Due in", "Expires in"))
                    else "red" if v in ("Overdue", "Failed", "Expired") else "grey")
        return None

    def write_table(ws, keys, headers, data):
        top, width = 1, max(len(headers), 4)
        if banner:
            for i, line in enumerate(banner, 1):
                ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=width)
                c = ws.cell(row=i, column=1, value=line)
                c.alignment = Alignment(horizontal="center")
                c.font = Font(bold=i in (1, 3), size=14 if i == 3 else 11,
                              color="B6202F" if i == 1 else "14203A")
            top = 6
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=top, column=j, value=h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
        for i, r in enumerate(data, top + 1):
            for j, k in enumerate(keys, 1):
                v = r.get(k)
                col = color(k, v)                       # colour from the original (English) value
                if v in (None, ""):
                    v = None
                else:
                    if k in DATE_KEYS and isinstance(v, str):
                        try:
                            v = date.fromisoformat(v)
                        except ValueError:
                            pass
                    v = export_value(k, v)
                cell = ws.cell(row=i, column=j, value=v)
                if isinstance(v, str) and v[:1] == "=":
                    cell.data_type = "s"                # text, never a formula
                if isinstance(v, date):
                    cell.number_format = "yyyy-mm-dd"
                if col:
                    cell.fill = fills[col]
        ws.freeze_panes = ws.cell(row=top + 1, column=1)
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{max(top + len(data), top)}"
        for j, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(ws.cell(row=i, column=j).value or ""))
                                           for i in range(top + 1, min(top + len(data), top + 200) + 1)])
            ws.column_dimensions[get_column_letter(j)].width = min(longest + 3, 42)

    keys, headers = [k for k, _ in cols], [h for _, h in cols]
    if per_station:
        groups = {}
        for r in rows:
            groups.setdefault(r["station"], []).append(r)
        if scope == "latest":
            summary = []
            for name, grp in groups.items():
                cnt = {"ok": 0, "soon": 0, "overdue": 0, "failed": 0, "never": 0}
                for r in grp:
                    cnt[status_key(r["status"])] += 1
                summary.append(dict(station=name, sensors=len(grp), **cnt))
            write_table(new_sheet(tr("Summary")),
                        ["station", "sensors", "ok", "soon", "overdue", "failed", "never"],
                        [tr(h) for h in ("Station", "Sensors", "OK", "Due within 30 days", "Overdue",
                                         "Failed", "Never calibrated")], summary)
        for name, grp in groups.items():
            write_table(new_sheet(name), keys, headers, grp)
    else:
        write_table(new_sheet(tr({"latest": "Register", "history": "History", "points": "Points"}[scope])),
                    keys, headers, rows)
    if with_standards:
        stds = [dict(x) for x in get_db().execute("SELECT * FROM reference_standards ORDER BY code")]
        for x in stds:
            x["status"] = standard_status(x)[0]
            x["active_txt"] = tr("Yes") if x["active"] else tr("No")
        write_table(new_sheet(tr("Reference standards")), [k for k, _ in STD_SHEET_COLUMNS],
                    [tr(h) for _, h in STD_SHEET_COLUMNS], stds)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@app.route("/export/download")
def export_download():
    rows, scope = export_rows(request.args)
    if not rows:
        flash("No records match those filters.")
        return redirect(url_for("export_csv"))
    chosen = set(request.args.getlist("cols"))
    cols = [c for c in EXPORT_COLUMNS if c[0] in chosen] or list(EXPORT_COLUMNS)
    if scope != "latest":                        # status only makes sense for the register view
        cols = [c for c in cols if c[0] != "status"]
    if scope != "points":                        # per-point columns only for the points export
        cols = [c for c in cols if c[0] not in ("point_no", "overall_result", "point_tolerance")]
        rename = {"reference_value": "Reference value (worst point)",
                  "measured_value": "Results (worst point)", "error": "Max error"}
        cols = [(k, rename.get(k, h)) for k, h in cols]
    cols = [(k, tr(h)) for k, h in cols]         # column headings follow the report language
    banner = None if request.args.get("nobanner") else BANNER[g.lang]
    fmt = request.args.get("format", "csv")
    stamp = f"calibration_{scope}_{export_label(request.args)}_{date.today().isoformat()}"
    if fmt == "xlsx":
        try:
            data = build_xlsx(rows, cols, scope, bool(request.args.get("per_station")),
                              bool(request.args.get("with_standards")), banner)
        except ImportError:
            flash("Could not create the Excel file: the openpyxl package is missing. Install it with "
                  "“python -m pip install openpyxl” and restart the app.")
            return redirect(url_for("export_csv"))
        return Response(data, mimetype=XLSX_MIME,
                        headers={"Content-Disposition": f'attachment; filename="{stamp}.xlsx"'})
    out = io.StringIO()
    w = csv.writer(out)
    if banner:
        for line in banner:
            w.writerow([line])
        w.writerow([])
    w.writerow([h for _, h in cols])
    for r in rows:
        w.writerow([safe_cell(export_value(k, r[k])) for k, _ in cols])
    data = out.getvalue()
    if fmt == "excel" or g.lang == "ne":         # UTF-8 BOM so Excel shows °C and Nepali text
        data = "\ufeff" + data
    return Response(data.encode("utf-8"), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{stamp}.csv"'})


#if __name__ == "__main__":
    # host="0.0.0.0" lets other PCs on your network connect; use "127.0.0.1" for this PC only
  #  app.run(host="127.0.0.1", port=5000, debug=True)